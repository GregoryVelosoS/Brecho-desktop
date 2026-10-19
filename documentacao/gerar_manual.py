"""Gera o manual didático do Sistema do Brechó em formato DOCX.

Dependência usada somente para gerar a documentação:
    python -m pip install python-docx
"""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


RAIZ_PROJETO = Path(__file__).resolve().parent.parent
PASTA_DOCUMENTACAO = Path(__file__).resolve().parent
ARQUIVO_SAIDA = PASTA_DOCUMENTACAO / "Manual_Projeto_Brecho_CustomTkinter.docx"

MARROM = "8B5E3C"
MARROM_CLARO = "EAD9CA"
BEGE = "F7F3EE"
CREME = "FFFDFC"
CINZA_CODIGO = "F2F2F2"
AZUL_NOTA = "DDEBF7"
VERDE_CHECK = "E2F0D9"
AMARELO_ATENCAO = "FFF2CC"
VERMELHO_ERRO = "FCE4D6"


def definir_sombreamento(celula, cor):
    propriedades = celula._tc.get_or_add_tcPr()
    sombreamento = propriedades.find(qn("w:shd"))
    if sombreamento is None:
        sombreamento = OxmlElement("w:shd")
        propriedades.append(sombreamento)
    sombreamento.set(qn("w:fill"), cor)


def definir_margens_celula(celula, superior=100, inicio=140, inferior=100, fim=140):
    propriedades = celula._tc.get_or_add_tcPr()
    margens = propriedades.first_child_found_in("w:tcMar")
    if margens is None:
        margens = OxmlElement("w:tcMar")
        propriedades.append(margens)

    for nome, valor in (
        ("top", superior),
        ("start", inicio),
        ("bottom", inferior),
        ("end", fim),
    ):
        margem = margens.find(qn(f"w:{nome}"))
        if margem is None:
            margem = OxmlElement(f"w:{nome}")
            margens.append(margem)
        margem.set(qn("w:w"), str(valor))
        margem.set(qn("w:type"), "dxa")


def impedir_quebra_linha_tabela(linha):
    propriedades = linha._tr.get_or_add_trPr()
    propriedades.append(OxmlElement("w:cantSplit"))


def adicionar_campo(paragrafo, instrucao):
    inicio = OxmlElement("w:fldChar")
    inicio.set(qn("w:fldCharType"), "begin")
    texto = OxmlElement("w:instrText")
    texto.set(qn("xml:space"), "preserve")
    texto.text = instrucao
    separador = OxmlElement("w:fldChar")
    separador.set(qn("w:fldCharType"), "separate")
    fim = OxmlElement("w:fldChar")
    fim.set(qn("w:fldCharType"), "end")
    execucao = paragrafo.add_run()
    execucao._r.extend((inicio, texto, separador, fim))


def configurar_documento(documento):
    secao = documento.sections[0]
    secao.page_height = Cm(29.7)
    secao.page_width = Cm(21)
    secao.top_margin = Cm(2)
    secao.bottom_margin = Cm(1.8)
    secao.left_margin = Cm(2.2)
    secao.right_margin = Cm(2.2)

    estilos = documento.styles
    normal = estilos["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    titulo = estilos["Title"]
    titulo.font.name = "Aptos Display"
    titulo.font.size = Pt(28)
    titulo.font.bold = True
    titulo.font.color.rgb = RGBColor.from_string(MARROM)

    for nome, tamanho in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 11.5)):
        estilo = estilos[nome]
        estilo.font.name = "Aptos Display"
        estilo.font.size = Pt(tamanho)
        estilo.font.bold = True
        estilo.font.color.rgb = RGBColor.from_string(MARROM)
        estilo.paragraph_format.space_before = Pt(10)
        estilo.paragraph_format.space_after = Pt(5)

    estilos["Heading 1"].paragraph_format.page_break_before = True

    propriedades = documento.core_properties
    propriedades.title = "Manual didático — Sistema do Brechó com Python e CustomTkinter"
    propriedades.subject = "Construção procedural, modular e comentada para iniciantes"
    propriedades.author = "Projeto Brechó — material didático"
    propriedades.keywords = "Python, CustomTkinter, MySQL, programação procedural, módulos"

    cabecalho = secao.header.paragraphs[0]
    cabecalho.text = "MANUAL DIDÁTICO  |  SISTEMA DO BRECHÓ"
    cabecalho.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    cabecalho.runs[0].font.name = "Aptos"
    cabecalho.runs[0].font.size = Pt(8)
    cabecalho.runs[0].font.color.rgb = RGBColor.from_string(MARROM)

    rodape = secao.footer.paragraphs[0]
    rodape.alignment = WD_ALIGN_PARAGRAPH.CENTER
    texto = rodape.add_run("Página ")
    texto.font.size = Pt(8)
    adicionar_campo(rodape, "PAGE")


def adicionar_titulo(documento, texto, subtitulo=None):
    paragrafo = documento.add_paragraph()
    paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragrafo.paragraph_format.space_before = Pt(110)
    execucao = paragrafo.add_run(texto)
    execucao.font.name = "Aptos Display"
    execucao.font.size = Pt(30)
    execucao.font.bold = True
    execucao.font.color.rgb = RGBColor.from_string(MARROM)

    if subtitulo:
        p_subtitulo = documento.add_paragraph()
        p_subtitulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_subtitulo.paragraph_format.space_before = Pt(10)
        trecho = p_subtitulo.add_run(subtitulo)
        trecho.font.name = "Aptos"
        trecho.font.size = Pt(14)
        trecho.font.color.rgb = RGBColor(70, 70, 70)


def adicionar_caixa(documento, titulo, texto, cor=AZUL_NOTA):
    tabela = documento.add_table(rows=1, cols=1)
    tabela.autofit = True
    celula = tabela.cell(0, 0)
    definir_sombreamento(celula, cor)
    definir_margens_celula(celula)
    impedir_quebra_linha_tabela(tabela.rows[0])
    paragrafo = celula.paragraphs[0]
    paragrafo.paragraph_format.space_after = Pt(0)
    parte_titulo = paragrafo.add_run(f"{titulo}: ")
    parte_titulo.bold = True
    paragrafo.add_run(texto)
    documento.add_paragraph().paragraph_format.space_after = Pt(0)


def adicionar_codigo(documento, codigo, legenda=None):
    if legenda:
        paragrafo_legenda = documento.add_paragraph()
        paragrafo_legenda.paragraph_format.space_after = Pt(3)
        trecho = paragrafo_legenda.add_run(legenda)
        trecho.bold = True
        trecho.font.size = Pt(9)
        trecho.font.color.rgb = RGBColor.from_string(MARROM)

    tabela = documento.add_table(rows=1, cols=1)
    tabela.autofit = True
    celula = tabela.cell(0, 0)
    celula.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    definir_sombreamento(celula, CINZA_CODIGO)
    definir_margens_celula(celula, superior=90, inicio=130, inferior=90, fim=130)

    paragrafo = celula.paragraphs[0]
    paragrafo.paragraph_format.space_after = Pt(0)
    paragrafo.paragraph_format.line_spacing = 1
    linhas = codigo.strip("\n").splitlines()
    for indice, linha in enumerate(linhas):
        trecho = paragrafo.add_run(linha)
        trecho.font.name = "Consolas"
        trecho.font.size = Pt(8)
        if indice < len(linhas) - 1:
            trecho.add_break()

    documento.add_paragraph().paragraph_format.space_after = Pt(0)


def adicionar_lista(documento, itens, numerada=False):
    estilo = "List Number" if numerada else "List Bullet"
    for item in itens:
        documento.add_paragraph(item, style=estilo)


def adicionar_tabela(documento, cabecalhos, linhas, larguras=None):
    tabela = documento.add_table(rows=1, cols=len(cabecalhos))
    tabela.style = "Table Grid"
    tabela.autofit = True
    for indice, cabecalho in enumerate(cabecalhos):
        celula = tabela.rows[0].cells[indice]
        celula.text = cabecalho
        definir_sombreamento(celula, MARROM_CLARO)
        for trecho in celula.paragraphs[0].runs:
            trecho.bold = True
        if larguras:
            celula.width = Cm(larguras[indice])

    for dados in linhas:
        celulas = tabela.add_row().cells
        for indice, valor in enumerate(dados):
            celulas[indice].text = str(valor)
            celulas[indice].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if larguras:
                celulas[indice].width = Cm(larguras[indice])
        impedir_quebra_linha_tabela(tabela.rows[-1])

    documento.add_paragraph().paragraph_format.space_after = Pt(0)
    return tabela


def adicionar_sumario(documento):
    documento.add_heading("Sumário", level=1)
    documento.add_paragraph(
        "No Microsoft Word, clique com o botão direito sobre o campo abaixo e "
        "escolha “Atualizar Campo” para preencher ou atualizar o sumário."
    )
    paragrafo = documento.add_paragraph()
    adicionar_campo(paragrafo, 'TOC \\o "1-3" \\h \\z \\u')


def ler_codigo(caminho_relativo):
    return (RAIZ_PROJETO / caminho_relativo).read_text(encoding="utf-8")


def adicionar_checkpoint(documento, itens):
    adicionar_caixa(
        documento,
        "Checkpoint",
        "Antes de continuar, confirme os itens abaixo.",
        VERDE_CHECK,
    )
    for item in itens:
        documento.add_paragraph(f"☐ {item}")


def gerar_manual():
    documento = Document()
    configurar_documento(documento)

    # Capa
    adicionar_titulo(
        documento,
        "Manual didático\nSistema do Brechó",
        "Python procedural, CustomTkinter e MySQL — construção modular do início ao fim",
    )
    documento.add_paragraph()
    caixa = documento.add_table(rows=1, cols=1)
    celula = caixa.cell(0, 0)
    definir_sombreamento(celula, BEGE)
    definir_margens_celula(celula, superior=180, inicio=220, inferior=180, fim=220)
    texto_capa = celula.paragraphs[0]
    texto_capa.alignment = WD_ALIGN_PARAGRAPH.CENTER
    texto_capa.add_run(
        "Público: estudantes iniciantes em programação\n"
        "Abordagem: procedural e modular\n"
        "Versão do manual: 1.0 — outubro de 2026"
    )
    documento.add_paragraph()
    aviso = documento.add_paragraph()
    aviso.alignment = WD_ALIGN_PARAGRAPH.CENTER
    aviso.add_run(
        "Objetivo: compreender cada etapa, executar o projeto e conseguir "
        "reaproveitar seus módulos em novos sistemas."
    ).italic = True

    adicionar_sumario(documento)

    # 1
    documento.add_heading("1. Apresentação do projeto", level=1)
    documento.add_paragraph(
        "Este manual ensina a construir uma aplicação desktop para um brechó. "
        "Ela possui login didático, menu principal, cadastro de clientes, cadastro "
        "de produtos e registro de pedidos com cálculo de descontos. Os dados são "
        "gravados em um banco MySQL."
    )
    documento.add_heading("1.1 O que você aprenderá", level=2)
    adicionar_lista(
        documento,
        [
            "Organizar um programa Python em módulos, isto é, arquivos com responsabilidades específicas.",
            "Criar funções, enviar dados por parâmetros e devolver valores com return.",
            "Construir janelas e formulários com CustomTkinter.",
            "Ligar botões a funções por meio de callbacks.",
            "Validar dados e tratar erros com try e except.",
            "Conectar Python ao MySQL e executar INSERT com parâmetros.",
            "Separar interface, regras de negócio, tema e acesso ao banco.",
            "Testar cada etapa antes de avançar para a seguinte.",
        ],
    )
    adicionar_caixa(
        documento,
        "Importante",
        "O código do projeto é procedural: não criamos classes próprias. "
        "Entretanto, o CustomTkinter é uma biblioteca orientada a objetos. Por "
        "isso aparecem chamadas como campo.get() e botao.pack(). Estamos usando "
        "objetos fornecidos pela biblioteca, sem exigir que o estudante crie classes.",
        AMARELO_ATENCAO,
    )
    documento.add_heading("1.2 Resultado final", level=2)
    adicionar_tabela(
        documento,
        ["Tela", "Responsabilidade", "Dados principais"],
        [
            ("Login", "Controlar a entrada didática", "Usuário e senha"),
            ("Menu", "Direcionar para os módulos", "Botões de navegação"),
            ("Clientes", "Cadastrar uma cliente", "Nome e tipo"),
            ("Produtos", "Cadastrar uma peça", "Nome, tamanho, categoria e preço"),
            ("Pedidos", "Calcular e salvar uma venda", "Cliente, produto, quantidade e descontos"),
        ],
        [3, 6, 7],
    )

    # 2
    documento.add_heading("2. Progressão recomendada para iniciantes", level=1)
    documento.add_paragraph(
        "A ordem abaixo reduz a quantidade de conceitos novos apresentados ao mesmo "
        "tempo. Primeiro entendemos dados e funções; depois construímos a interface; "
        "por último conectamos tudo."
    )
    adicionar_tabela(
        documento,
        ["Etapa", "Foco", "Evidência de aprendizagem"],
        [
            ("1", "Ambiente e execução", "Conseguir executar um arquivo .py"),
            ("2", "Funções e módulos", "Importar e chamar uma função simples"),
            ("3", "Banco e tabelas", "Criar o banco brecho_uc2"),
            ("4", "Regras de negócio", "Calcular um pedido sem abrir janela"),
            ("5", "CustomTkinter básico", "Criar janela, texto, campo e botão"),
            ("6", "Componentes reutilizáveis", "Aplicar o mesmo estilo em várias telas"),
            ("7", "Telas separadas", "Navegar entre os módulos"),
            ("8", "Integração MySQL", "Salvar e consultar registros"),
            ("9", "Testes e evolução", "Identificar erros e adicionar uma melhoria"),
        ],
        [2, 6, 8],
    )
    adicionar_caixa(
        documento,
        "Estratégia de estudo",
        "Em cada etapa use o ciclo ação → observação → explicação → nova ação. "
        "Digite o código, execute, observe o resultado, explique com suas palavras "
        "e então faça uma pequena alteração.",
    )

    # 3
    documento.add_heading("3. Preparação do ambiente", level=1)
    documento.add_heading("3.1 Programas necessários", level=2)
    adicionar_lista(
        documento,
        [
            "Python 3: interpreta e executa os arquivos .py.",
            "Visual Studio Code ou outro editor: permite organizar e editar o projeto.",
            "MySQL Server: armazena os dados.",
            "MySQL Workbench: facilita a execução do script SQL e a consulta das tabelas.",
        ],
    )
    documento.add_heading("3.2 Criar e ativar um ambiente virtual", level=2)
    documento.add_paragraph(
        "O ambiente virtual mantém as bibliotecas deste projeto separadas das de "
        "outros programas. Isso diminui conflitos de versões."
    )
    adicionar_codigo(
        documento,
        """python -m venv .venv
.venv\\Scripts\\activate""",
        "Comandos no PowerShell, dentro da pasta do projeto",
    )
    adicionar_caixa(
        documento,
        "Se o PowerShell bloquear a ativação",
        "Use o Prompt de Comando com .venv\\Scripts\\activate.bat ou consulte a "
        "política de execução definida pela instituição. Não desative proteções do "
        "computador sem autorização.",
        AMARELO_ATENCAO,
    )
    documento.add_heading("3.3 Instalar as bibliotecas", level=2)
    adicionar_codigo(documento, "python -m pip install -r requirements.txt")
    adicionar_tabela(
        documento,
        ["Biblioteca", "Versão do projeto", "Função"],
        [
            ("customtkinter", "5.2.2", "Fornecer componentes gráficos com aparência moderna"),
            ("mysql-connector-python", "9.0.0", "Conectar o programa ao servidor MySQL"),
            ("tkinter", "Incluído no Python", "Fornecer caixas de mensagem e a base gráfica"),
        ],
        [5, 3, 8],
    )
    adicionar_checkpoint(
        documento,
        [
            "python --version mostra uma versão do Python 3.",
            "pip install -r requirements.txt termina sem erro.",
            "O MySQL Server está iniciado.",
        ],
    )

    # 4
    documento.add_heading("4. Planejamento e estrutura modular", level=1)
    documento.add_paragraph(
        "Antes de programar, separamos responsabilidades. Uma tela não deve conhecer "
        "detalhes de cores, fórmulas e conexão ao banco ao mesmo tempo. Essa divisão "
        "facilita manutenção, testes e reaproveitamento."
    )
    adicionar_codigo(
        documento,
        """Brecho-desktop/
├── app.py                    # inicia o programa e controla a navegação
├── config.py                 # dados locais de conexão
├── database.py               # comandos enviados ao MySQL
├── regras_negocio.py         # cálculos independentes da interface
├── tema.py                   # paleta visual
├── requirements.txt          # bibliotecas externas
├── sql/
│   └── schema.sql            # criação do banco e das tabelas
└── telas/
    ├── __init__.py           # identifica a pasta como pacote Python
    ├── componentes.py        # componentes visuais padronizados
    ├── login.py
    ├── menu.py
    ├── clientes.py
    ├── produtos.py
    └── pedidos.py""",
    )
    documento.add_heading("4.1 Fluxo de dependências", level=2)
    adicionar_codigo(
        documento,
        """app.py
  ├─ chama as funções criar_tela_*()
  ├─ telas usam telas/componentes.py
  │      └─ componentes.py usa tema.py
  ├─ telas usam database.py
  │      └─ database.py usa config.py
  └─ pedidos.py usa regras_negocio.py""",
    )
    adicionar_caixa(
        documento,
        "Por que evitar importação circular",
        "Se app.py importa uma tela e essa tela importa app.py de volta, os arquivos "
        "podem ser carregados pela metade. Por isso app.py envia as funções de "
        "navegação por meio de um dicionário.",
    )

    # 5
    documento.add_heading("5. Fundamentos de Python usados no projeto", level=1)
    documento.add_heading("5.1 Funções, parâmetros e return", level=2)
    adicionar_codigo(
        documento,
        """def somar(valor_a, valor_b):
    resultado = valor_a + valor_b
    return resultado

total = somar(10, 5)""",
    )
    documento.add_paragraph(
        "def inicia a definição; valor_a e valor_b são parâmetros; return devolve o "
        "resultado. No projeto, esse padrão é usado nos cálculos e nas funções que "
        "criam telas."
    )
    documento.add_heading("5.2 Módulos e importações", level=2)
    adicionar_codigo(
        documento,
        """from regras_negocio import calcular_pedido

subtotal, desconto, total = calcular_pedido(100, 2, 10, "VIP")""",
    )
    documento.add_paragraph(
        "Um módulo é um arquivo Python. A importação permite usar uma função de outro "
        "arquivo sem copiá-la. A documentação oficial do Python recomenda módulos "
        "para organizar programas maiores e reutilizar definições."
    )
    documento.add_heading("5.3 Dicionários", level=2)
    adicionar_codigo(
        documento,
        """campos = {
    "nome": campo_nome,
    "preco": campo_preco,
}

nome_digitado = campos["nome"].get()""",
    )
    documento.add_paragraph(
        "O dicionário relaciona uma chave a um valor. Agrupar componentes em um "
        "dicionário evita funções com uma lista muito grande de parâmetros."
    )
    documento.add_heading("5.4 Tratamento de erros", level=2)
    adicionar_codigo(
        documento,
        """try:
    preco = float(campo_preco.get())
except ValueError:
    messagebox.showerror("Erro", "Digite um preço numérico.")""",
    )
    documento.add_paragraph(
        "try contém uma operação que pode falhar. except trata um erro conhecido. "
        "Capturar erros específicos torna o diagnóstico mais claro do que capturar "
        "qualquer Exception."
    )

    # 6
    documento.add_heading("6. Criando o banco de dados", level=1)
    documento.add_paragraph(
        "Abra o MySQL Workbench, crie uma nova aba SQL, copie o conteúdo de "
        "sql/schema.sql e execute o script completo."
    )
    adicionar_codigo(
        documento,
        """CREATE DATABASE IF NOT EXISTS brecho_uc2;
USE brecho_uc2;

CREATE TABLE IF NOT EXISTS clientes (
    id_cliente INT AUTO_INCREMENT PRIMARY KEY,
    nome VARCHAR(100) NOT NULL,
    tipo VARCHAR(20) NOT NULL
);""",
        "Primeira tabela do projeto",
    )
    adicionar_tabela(
        documento,
        ["Termo SQL", "Significado"],
        [
            ("CREATE DATABASE", "Cria o banco que agrupará as tabelas"),
            ("USE", "Seleciona o banco para os próximos comandos"),
            ("CREATE TABLE", "Cria uma estrutura de linhas e colunas"),
            ("AUTO_INCREMENT", "Gera o próximo código automaticamente"),
            ("PRIMARY KEY", "Identifica cada registro de forma única"),
            ("NOT NULL", "Torna o preenchimento obrigatório"),
            ("DECIMAL(10,2)", "Guarda valores monetários com duas casas decimais"),
        ],
        [5, 11],
    )
    documento.add_heading("6.1 Modelo adotado", level=2)
    documento.add_paragraph(
        "Nesta versão didática, pedidos armazenam cliente e produto como texto. Isso "
        "evita introduzir chaves estrangeiras e itens de pedido antes de o estudante "
        "dominar o fluxo básico. Em uma evolução, o recomendado é relacionar IDs e "
        "criar uma tabela itens_pedido."
    )
    adicionar_checkpoint(
        documento,
        [
            "O banco brecho_uc2 aparece no painel SCHEMAS.",
            "As tabelas clientes, produtos e pedidos foram criadas.",
            "SHOW TABLES; apresenta as três tabelas.",
        ],
    )

    # 7
    documento.add_heading("7. Configurando a conexão", level=1)
    documento.add_paragraph(
        "O arquivo config.py guarda as informações específicas do computador. Use a "
        "senha do usuário MySQL local."
    )
    adicionar_codigo(
        documento,
        """DB_CONFIG = {
    "host": "localhost",
    "user": "root",
    "password": "SUA_SENHA_AQUI",
    "database": "brecho_uc2",
}""",
    )
    adicionar_caixa(
        documento,
        "Segurança",
        "Este formato é aceitável apenas para um protótipo local e didático. Não "
        "publique senhas reais em repositórios. Em projetos distribuídos, use "
        "variáveis de ambiente e um arquivo .env ignorado pelo Git.",
        AMARELO_ATENCAO,
    )

    # 8
    documento.add_heading("8. Implementando o acesso ao MySQL", level=1)
    documento.add_paragraph(
        "database.py concentra a conexão e os INSERTs. As telas apenas informam os "
        "dados; não precisam conhecer SQL."
    )
    adicionar_codigo(
        documento,
        """def conectar():
    try:
        return mysql.connector.connect(**DB_CONFIG)
    except Error as erro:
        raise ConnectionError(
            f"Não foi possível conectar ao MySQL: {erro}"
        ) from erro""",
        "Função de conexão",
    )
    documento.add_paragraph(
        "Os dois asteriscos transformam o dicionário DB_CONFIG em argumentos "
        "nomeados. Se a conexão falhar, o erro da biblioteca é transformado em uma "
        "mensagem que a tela consegue apresentar."
    )
    adicionar_codigo(
        documento,
        """sql = "INSERT INTO clientes (nome, tipo) VALUES (%s, %s)"
cursor.execute(sql, (nome, tipo))
conexao.commit()""",
        "Consulta parametrizada",
    )
    adicionar_caixa(
        documento,
        "Por que usar %s",
        "Os valores são enviados separadamente do comando SQL. O conector converte "
        "os tipos e evita concatenar diretamente o texto digitado pelo usuário, o "
        "que é mais seguro e menos sujeito a erros de aspas.",
    )
    adicionar_tabela(
        documento,
        ["Operação", "Papel"],
        [
            ("conexao.cursor()", "Cria o objeto que envia comandos SQL"),
            ("cursor.execute()", "Executa o comando com seus parâmetros"),
            ("conexao.commit()", "Confirma uma alteração"),
            ("conexao.rollback()", "Desfaz a transação após um erro"),
            ("cursor.lastrowid", "Obtém o ID AUTO_INCREMENT criado"),
            ("cursor.close()", "Libera o cursor"),
            ("conexao.close()", "Encerra a conexão"),
        ],
        [6, 10],
    )

    # 9
    documento.add_heading("9. Regras de negócio independentes da tela", level=1)
    documento.add_paragraph(
        "Uma regra de negócio descreve como o sistema calcula ou decide algo. Ela "
        "deve funcionar mesmo sem janela e sem banco, facilitando testes."
    )
    adicionar_codigo(
        documento,
        """def calcular_valor_produto(preco, quantidade, desconto_produto):
    valor_bruto = preco * quantidade
    valor_desconto = valor_bruto * desconto_produto / 100
    return valor_bruto - valor_desconto""",
    )
    adicionar_codigo(
        documento,
        """def obter_desconto_cliente(tipo):
    if tipo == "Novo":
        return 5
    if tipo == "VIP":
        return 15
    return 0""",
    )
    documento.add_heading("9.1 Exemplo acompanhado", level=2)
    adicionar_tabela(
        documento,
        ["Passo", "Conta", "Resultado"],
        [
            ("Valor bruto", "R$ 100 × 2", "R$ 200"),
            ("Desconto do produto", "10% de R$ 200", "R$ 20"),
            ("Subtotal", "R$ 200 − R$ 20", "R$ 180"),
            ("Desconto VIP", "15% de R$ 180", "R$ 27"),
            ("Valor final", "R$ 180 − R$ 27", "R$ 153"),
        ],
        [4, 7, 5],
    )
    adicionar_codigo(
        documento,
        """python -c "from regras_negocio import calcular_pedido; print(calcular_pedido(100, 2, 10, 'VIP'))"

# Resultado esperado:
# (180.0, 15, 153.0)""",
        "Teste rápido no terminal",
    )

    # 10
    documento.add_heading("10. Tema visual centralizado", level=1)
    documento.add_paragraph(
        "tema.py contém a paleta visual. Centralizar cores impede que cada tela repita "
        "muitas importações e permite mudar a aparência em um único lugar."
    )
    adicionar_codigo(
        documento,
        """COR_FUNDO = "#F7F3EE"
COR_CARD = "#FFFFFF"
COR_PRINCIPAL = "#8B5E3C"
COR_HOVER = "#6F472D"
COR_DESTAQUE = "#C9A27E"
COR_TEXTO = "#2D2A26"
COR_TEXTO_CLARO = "#FFFFFF"
COR_BORDA = "#D8C3B0"
COR_CAMPO = "#FFFDFC"
""",
    )
    adicionar_codigo(
        documento,
        """def configurar_tema():
    ctk.set_appearance_mode("light")
    ctk.set_default_color_theme("blue")""",
    )
    documento.add_paragraph(
        "O modo de aparência deve ser definido antes de criar a janela. As cores "
        "próprias substituem a maior parte do tema-base nos componentes."
    )

    # 11
    documento.add_heading("11. Fundamentos do CustomTkinter", level=1)
    documento.add_paragraph(
        "CustomTkinter fornece componentes gráficos semelhantes aos do Tkinter, com "
        "mais opções de aparência. Todo componente precisa de um elemento pai. A "
        "janela contém telas; telas contêm cards; cards contêm textos, campos e botões."
    )
    documento.add_heading("11.1 Principais componentes", level=2)
    adicionar_tabela(
        documento,
        ["Componente", "Uso no projeto", "Exemplo"],
        [
            ("CTk", "Janela principal", "janela = ctk.CTk()"),
            ("CTkFrame", "Tela, card ou área de organização", "ctk.CTkFrame(pai)"),
            ("CTkLabel", "Títulos, instruções e resultados", "ctk.CTkLabel(pai, text='...')"),
            ("CTkEntry", "Entrada de texto", "campo.get()"),
            ("CTkButton", "Executar uma ação", "command=salvar"),
            ("CTkComboBox", "Selecionar uma opção", "combo.set('Novo')"),
            ("CTkFont", "Definir tamanho e peso", "ctk.CTkFont(size=18, weight='bold')"),
        ],
        [4, 7, 6],
    )
    documento.add_heading("11.2 Gerenciadores de posição", level=2)
    adicionar_tabela(
        documento,
        ["Método", "Quando usar", "Conceito"],
        [
            ("pack()", "Blocos empilhados", "Organiza por lados e preenchimento"),
            ("grid()", "Formulários", "Organiza em linhas e colunas"),
            ("place()", "Posição relativa ou exata", "Usado no card central do login"),
        ],
        [3, 6, 8],
    )
    adicionar_caixa(
        documento,
        "Erro comum",
        "Não misture pack() e grid() para organizar filhos que possuem o mesmo pai. "
        "É permitido usar pack em uma tela e grid dentro de um card, pois são pais "
        "diferentes.",
        VERMELHO_ERRO,
    )
    documento.add_heading("11.3 Parâmetros usados com frequência", level=2)
    adicionar_tabela(
        documento,
        ["Parâmetro", "Explicação"],
        [
            ("text", "Texto apresentado"),
            ("fg_color", "Cor de fundo"),
            ("text_color", "Cor do texto"),
            ("hover_color", "Cor do botão ao passar o mouse"),
            ("command", "Função executada no clique; deve ser enviada sem parênteses"),
            ("values", "Lista de opções de um CTkComboBox"),
            ("placeholder_text", "Texto de orientação dentro de um campo vazio"),
            ("show='*'", "Oculta visualmente uma senha"),
            ("padx / pady", "Margens horizontal e vertical"),
            ("sticky='ew'", "Estica o componente nas direções leste e oeste"),
            ("columnspan", "Faz o componente ocupar mais de uma coluna"),
        ],
        [5, 12],
    )
    documento.add_heading("11.4 Métodos importantes", level=2)
    adicionar_tabela(
        documento,
        ["Método", "Efeito"],
        [
            ("get()", "Lê o valor atual de um campo ou combo"),
            ("delete(0, 'end')", "Apaga o texto de um campo"),
            ("set(valor)", "Define a seleção de um combo"),
            ("configure(...) ", "Altera um componente já criado"),
            ("bind(evento, função)", "Liga um evento de teclado ou mouse a uma função"),
            ("destroy()", "Remove e encerra um componente"),
            ("mainloop()", "Mantém a janela processando eventos"),
        ],
        [6, 11],
    )

    # 12
    documento.add_heading("12. Criando componentes reutilizáveis", level=1)
    documento.add_paragraph(
        "telas/componentes.py é a camada visual comum. As outras telas importam o "
        "módulo como ui e não precisam importar todas as cores."
    )
    adicionar_codigo(
        documento,
        """from telas import componentes as ui

tela = ui.criar_tela(janela)
campo_nome = ui.criar_campo(tela)
botao = ui.criar_botao(tela, "Salvar", salvar)""",
    )
    adicionar_codigo(
        documento,
        """def criar_botao(elemento_pai, texto, comando, largura=None, secundario=False):
    cor_fundo = tema.COR_DESTAQUE if secundario else tema.COR_PRINCIPAL
    configuracao = {
        "text": texto,
        "fg_color": cor_fundo,
        "hover_color": tema.COR_HOVER,
        "text_color": tema.COR_TEXTO_CLARO,
        "command": comando,
    }

    if largura is not None:
        configuracao["width"] = largura

    return ctk.CTkButton(elemento_pai, **configuracao)""",
        "Exemplo de fábrica procedural de componentes",
    )
    documento.add_paragraph(
        "A função recebe apenas o que muda e aplica internamente o estilo que se "
        "repete. O dicionário configuracao é expandido com ** ao chamar CTkButton."
    )

    # 13
    documento.add_heading("13. Janela principal e navegação", level=1)
    documento.add_paragraph(
        "app.py controla uma única janela. Para trocar de tela, destrói o frame atual "
        "e solicita que outra função crie o próximo."
    )
    adicionar_codigo(
        documento,
        """janela = None
frame_atual = None

def trocar_tela(funcao_criadora):
    global frame_atual

    if frame_atual is not None:
        frame_atual.destroy()

    frame_atual = funcao_criadora(janela, navegacao)
    frame_atual.pack(fill="both", expand=True)""",
    )
    documento.add_paragraph(
        "global informa que a função atualizará a variável definida fora dela. O "
        "parâmetro funcao_criadora recebe uma função como valor. Isso permite usar a "
        "mesma lógica para todas as telas."
    )
    adicionar_codigo(
        documento,
        """navegacao = {
    "login": abrir_login,
    "menu": abrir_menu,
    "clientes": abrir_clientes,
    "produtos": abrir_produtos,
    "pedidos": abrir_pedidos,
}""",
    )
    adicionar_caixa(
        documento,
        "Atenção ao command",
        "Use command=abrir_menu, sem parênteses. command=abrir_menu() executaria a "
        "função imediatamente durante a criação do botão.",
        AMARELO_ATENCAO,
    )

    # 14
    documento.add_heading("14. Construindo a tela de login", level=1)
    adicionar_codigo(
        documento,
        """def validar_login(campo_usuario, campo_senha, abrir_menu):
    usuario_digitado = campo_usuario.get()
    senha_digitada = campo_senha.get()

    if usuario_digitado == USUARIO_CORRETO and senha_digitada == SENHA_CORRETA:
        abrir_menu()
        return

    messagebox.showerror(
        "Acesso negado",
        "Usuário ou senha incorretos.",
    )""",
    )
    documento.add_paragraph(
        "A função recebe os componentes e a função de navegação. get() lê os textos. "
        "return encerra a função após um login válido."
    )
    adicionar_codigo(
        documento,
        """def tentar_entrar():
    validar_login(
        campo_usuario,
        campo_senha,
        navegacao["menu"],
    )

ui.criar_botao(
    card,
    texto="Entrar",
    comando=tentar_entrar,
).pack()""",
        "Callback do botão",
    )
    documento.add_paragraph(
        "tentar_entrar é uma função interna: ela existe somente durante a construção "
        "desta tela e consegue acessar os campos locais."
    )

    # 15
    documento.add_heading("15. Construindo o menu", level=1)
    documento.add_paragraph(
        "O menu demonstra reutilização. Em vez de repetir todo o código de três cards, "
        "criamos a função criar_card_opcao() e mudamos apenas textos e comandos."
    )
    adicionar_codigo(
        documento,
        """criar_card_opcao(
    area_cards,
    0,
    "Clientes",
    "Cadastre clientes do brechó.",
    "Cadastrar clientes",
    navegacao["clientes"],
)""",
    )
    documento.add_paragraph(
        "O número 0 é a coluna do primeiro card. Os outros usam 1 e 2. Todos recebem "
        "weight=1 para dividir a largura disponível."
    )

    # 16
    documento.add_heading("16. Cadastro de clientes", level=1)
    adicionar_codigo(
        documento,
        """def salvar_cliente(campo_nome, campo_tipo):
    nome = campo_nome.get().strip()
    tipo = campo_tipo.get()

    if not nome:
        messagebox.showwarning(
            "Dados incompletos",
            "Informe o nome da cliente.",
        )
        return

    codigo = inserir_cliente(nome, tipo)""",
    )
    documento.add_paragraph(
        "strip() remove espaços no início e no fim. A condição if not nome rejeita "
        "texto vazio. Depois da validação, a função de banco recebe apenas valores, "
        "não os componentes gráficos."
    )
    adicionar_caixa(
        documento,
        "Separação importante",
        "A tela conhece campos e mensagens. database.py conhece SQL. Essa separação "
        "permite substituir a interface sem reescrever as consultas.",
    )

    # 17
    documento.add_heading("17. Cadastro de produtos", level=1)
    adicionar_codigo(
        documento,
        """campos = {
    "nome": ui.criar_campo_formulario(corpo, "Nome do produto", 0, 0),
    "tamanho": ui.criar_campo_formulario(corpo, "Tamanho", 0, 1),
}

preco = float(campos["preco"].get())
if preco < 0:
    raise ValueError("O preço não pode ser negativo.")""",
    )
    documento.add_paragraph(
        "Os campos ficam agrupados por nome. float() converte o texto para um número "
        "decimal. Se a conversão falhar ou o valor for negativo, a tela apresenta uma "
        "mensagem e não envia o INSERT."
    )
    adicionar_checkpoint(
        documento,
        [
            "Nome vazio produz um aviso.",
            "Preço com letras produz uma mensagem de dado inválido.",
            "Preço negativo é rejeitado.",
            "Um produto válido recebe um código do MySQL.",
        ],
    )

    # 18
    documento.add_heading("18. Cadastro e cálculo de pedidos", level=1)
    documento.add_paragraph(
        "Pedidos combinam interface, validação, regra de negócio e banco. Por isso o "
        "módulo é dividido em quatro funções: ler_numeros, calcular_valores_pedido, "
        "salvar_pedido e limpar_pedido."
    )
    adicionar_codigo(
        documento,
        """def ler_numeros(campos):
    quantidade = int(campos["quantidade"].get())
    preco = float(campos["preco"].get())
    desconto_produto = float(campos["desconto_produto"].get())

    if quantidade <= 0:
        raise ValueError("A quantidade deve ser maior que zero.")
    if not 0 <= desconto_produto <= 100:
        raise ValueError("O desconto deve estar entre 0 e 100.")

    return quantidade, preco, desconto_produto""",
    )
    documento.add_heading("18.1 Estado temporário sem classes", level=2)
    adicionar_codigo(
        documento,
        """estado = {"resultado": None}

# Depois de calcular:
estado["resultado"] = (subtotal, desconto_cliente, valor_final)""",
    )
    documento.add_paragraph(
        "O dicionário estado guarda o último cálculo. Como dicionários são mutáveis, "
        "as funções podem atualizar o conteúdo sem usar self ou variável global."
    )
    documento.add_heading("18.2 Atualizando a interface", level=2)
    adicionar_codigo(
        documento,
        """resumo["subtotal"].configure(
    text=f"Subtotal: R$ {subtotal:.2f}"
)
resumo["total"].configure(
    text=f"Valor final: R$ {valor_final:.2f}"
)""",
    )
    documento.add_paragraph(
        "configure altera o texto de um componente já existente. :.2f formata o "
        "número com duas casas decimais."
    )

    # 19
    documento.add_heading("19. Executando o projeto completo", level=1)
    adicionar_lista(
        documento,
        [
            "Inicie o MySQL Server.",
            "Execute sql/schema.sql no MySQL Workbench.",
            "Ajuste config.py com os dados do banco local.",
            "Ative o ambiente virtual.",
            "Instale requirements.txt.",
            "Execute python app.py.",
            "Entre com admin / 1234.",
            "Cadastre um cliente e um produto.",
            "Preencha um pedido, clique em Calcular e depois em Salvar pedido.",
        ],
        numerada=True,
    )
    adicionar_codigo(
        documento,
        """USE brecho_uc2;

SELECT * FROM clientes;
SELECT * FROM produtos;
SELECT * FROM pedidos;""",
        "Conferência no MySQL Workbench",
    )
    adicionar_caixa(
        documento,
        "Critério de conclusão",
        "O projeto está integrado quando o registro exibido pelo SELECT corresponde "
        "ao que foi informado na interface e o valor_final corresponde ao cálculo.",
        VERDE_CHECK,
    )

    # 20
    documento.add_heading("20. Testes e depuração", level=1)
    documento.add_heading("20.1 Testes manuais mínimos", level=2)
    adicionar_tabela(
        documento,
        ["Caso", "Entrada", "Resultado esperado"],
        [
            ("Login correto", "admin / 1234", "Abre o menu"),
            ("Login incorreto", "admin / senha errada", "Mostra acesso negado"),
            ("Cliente vazio", "Nome em branco", "Mostra aviso e não grava"),
            ("Produto inválido", "Preço abc", "Mostra erro numérico"),
            ("Pedido normal", "100 × 2, 10%, VIP", "Subtotal 180 e total 153"),
            ("Quantidade zero", "Quantidade 0", "Mostra erro"),
            ("Desconto acima de 100", "Desconto 101", "Mostra erro"),
            ("Banco parado", "Salvar com MySQL desligado", "Mostra erro de conexão"),
        ],
        [4, 6, 7],
    )
    documento.add_heading("20.2 Como ler uma mensagem de erro", level=2)
    adicionar_lista(
        documento,
        [
            "Leia a última linha primeiro: ela informa o tipo e a mensagem do erro.",
            "Procure o primeiro arquivo do seu projeto citado no traceback.",
            "Confira o número da linha e os valores usados naquele ponto.",
            "Corrija uma causa por vez e execute novamente.",
            "Não esconda erros desconhecidos com except Exception sem diagnóstico.",
        ],
    )
    documento.add_heading("20.3 Erros frequentes", level=2)
    adicionar_tabela(
        documento,
        ["Mensagem ou sintoma", "Causa provável", "Ação"],
        [
            ("ModuleNotFoundError", "Biblioteca não instalada ou terminal fora do ambiente", "Ativar .venv e executar pip install -r requirements.txt"),
            ("Access denied for user", "Usuário ou senha do MySQL incorretos", "Revisar config.py"),
            ("Unknown database", "schema.sql não foi executado", "Criar brecho_uc2 no Workbench"),
            ("Can't connect", "MySQL Server parado ou porta incorreta", "Iniciar o serviço e conferir host"),
            ("invalid command name", "Callback tentou usar componente destruído", "Revisar a troca de telas"),
            ("Janela abre e fecha", "mainloop não foi chamado ou ocorreu erro", "Executar pelo terminal e ler o traceback"),
        ],
        [5, 6, 7],
    )

    # 21
    documento.add_heading("21. Como reaproveitar e evoluir o projeto", level=1)
    documento.add_heading("21.1 Criar um novo módulo", level=2)
    adicionar_lista(
        documento,
        [
            "Crie telas/fornecedores.py.",
            "Implemente criar_tela_fornecedores(janela, navegacao).",
            "Use ui.criar_tela, ui.criar_cabecalho, ui.criar_campo e ui.criar_botao.",
            "Adicione inserir_fornecedor() em database.py e a tabela no schema.sql.",
            "Importe a função em app.py.",
            "Crie abrir_fornecedores() e registre a função no dicionário navegacao.",
            "Adicione um card no menu e teste dados válidos e inválidos.",
        ],
        numerada=True,
    )
    adicionar_codigo(
        documento,
        """from telas import componentes as ui

def criar_tela_fornecedores(janela, navegacao):
    tela = ui.criar_tela(janela)
    ui.criar_cabecalho(
        tela,
        "Cadastro de fornecedores",
        navegacao["menu"],
    )

    # Adicione aqui os campos e o botão.
    return tela""",
        "Modelo inicial de um novo módulo",
    )
    documento.add_heading("21.2 Evoluções sugeridas", level=2)
    adicionar_lista(
        documento,
        [
            "Listar, editar e excluir clientes e produtos.",
            "Buscar clientes e produtos cadastrados em vez de digitar nomes no pedido.",
            "Criar chaves estrangeiras e uma tabela itens_pedido.",
            "Usar Decimal para cálculos monetários.",
            "Mover credenciais para variáveis de ambiente.",
            "Criar testes automatizados para regras_negocio.py.",
            "Validar alterações nos campos depois do cálculo do pedido.",
        ],
    )

    # 22
    documento.add_heading("22. Referência rápida", level=1)
    adicionar_tabela(
        documento,
        ["Função do projeto", "Arquivo", "Responsabilidade"],
        [
            ("iniciar_aplicacao", "app.py", "Criar a janela e iniciar mainloop"),
            ("trocar_tela", "app.py", "Destruir a tela atual e montar outra"),
            ("configurar_tema", "tema.py", "Definir aparência antes da janela"),
            ("criar_botao", "telas/componentes.py", "Padronizar botões"),
            ("criar_campo_formulario", "telas/componentes.py", "Criar rótulo e entrada"),
            ("calcular_pedido", "regras_negocio.py", "Aplicar os descontos"),
            ("conectar", "database.py", "Abrir conexão MySQL"),
            ("inserir_cliente", "database.py", "Gravar cliente"),
            ("inserir_produto", "database.py", "Gravar produto"),
            ("inserir_pedido", "database.py", "Gravar pedido"),
        ],
        [5, 5, 8],
    )
    documento.add_heading("22.1 Glossário", level=2)
    adicionar_tabela(
        documento,
        ["Termo", "Definição"],
        [
            ("Callback", "Função guardada para ser executada após um evento"),
            ("Componente/widget", "Elemento visual como botão, campo ou texto"),
            ("Frame", "Contêiner que agrupa outros componentes"),
            ("Módulo", "Arquivo Python importável"),
            ("Pacote", "Pasta de módulos Python, normalmente com __init__.py"),
            ("Regra de negócio", "Cálculo ou decisão que representa o funcionamento do domínio"),
            ("Validação", "Verificação que impede dados incorretos"),
            ("Transação", "Conjunto de operações confirmado por commit ou desfeito por rollback"),
            ("SQL parametrizado", "Comando que envia os valores separadamente da instrução SQL"),
        ],
        [5, 13],
    )

    # 23
    documento.add_heading("23. Documentação oficial consultada", level=1)
    documento.add_paragraph(
        "Use estas fontes quando quiser confirmar parâmetros, métodos ou aprofundar um tema:"
    )
    adicionar_lista(
        documento,
        [
            "CustomTkinter — documentação oficial: https://customtkinter.tomschimansky.com/documentation/",
            "Python — módulos: https://docs.python.org/3/tutorial/modules.html",
            "Python — erros e exceções: https://docs.python.org/3/tutorial/errors.html",
            "Python — interface Tkinter: https://docs.python.org/3/library/tkinter.html",
            "MySQL Connector/Python — guia oficial: https://dev.mysql.com/doc/connector-python/en/",
            "MySQL — inserção e transações: https://dev.mysql.com/doc/connector-python/en/connector-python-example-cursor-transaction.html",
            "MySQLCursor.execute(): https://dev.mysql.com/doc/connector-python/en/connector-python-api-mysqlcursor-execute.html",
        ],
    )
    adicionar_caixa(
        documento,
        "Versões",
        "Este manual acompanha customtkinter 5.2.2 e mysql-connector-python 9.0.0, "
        "conforme requirements.txt. Documentações on-line podem evoluir.",
    )

    # Apêndice de código
    documento.add_heading("Apêndice A — Código completo por módulo", level=1)
    documento.add_paragraph(
        "Os blocos seguintes correspondem à versão procedural usada na elaboração "
        "deste manual. Compare com os arquivos do projeto após concluir cada etapa."
    )

    arquivos = [
        ("requirements.txt", "requirements.txt"),
        ("sql/schema.sql", "sql/schema.sql"),
        ("database.py", "database.py"),
        ("regras_negocio.py", "regras_negocio.py"),
        ("tema.py", "tema.py"),
        ("app.py", "app.py"),
        ("telas/__init__.py", "telas/__init__.py"),
        ("telas/componentes.py", "telas/componentes.py"),
        ("telas/login.py", "telas/login.py"),
        ("telas/menu.py", "telas/menu.py"),
        ("telas/clientes.py", "telas/clientes.py"),
        ("telas/produtos.py", "telas/produtos.py"),
        ("telas/pedidos.py", "telas/pedidos.py"),
    ]
    for titulo_arquivo, caminho in arquivos:
        documento.add_heading(titulo_arquivo, level=2)
        adicionar_codigo(documento, ler_codigo(caminho))

    documento.add_heading("config.py — modelo seguro para copiar", level=2)
    adicionar_codigo(
        documento,
        """DB_CONFIG = {
    "host": "localhost",
    "user": "root",
    "password": "SUA_SENHA_AQUI",
    "database": "brecho_uc2",
}""",
    )
    adicionar_caixa(
        documento,
        "Fim do percurso",
        "Você concluiu o manual quando consegue explicar o caminho de um clique no "
        "botão Salvar até o INSERT no MySQL e consegue criar um novo módulo seguindo "
        "o mesmo padrão.",
        VERDE_CHECK,
    )

    # Solicita ao Word que atualize campos como sumário ao abrir o arquivo.
    configuracoes = documento.settings.element
    atualizar_campos = OxmlElement("w:updateFields")
    atualizar_campos.set(qn("w:val"), "true")
    configuracoes.append(atualizar_campos)

    PASTA_DOCUMENTACAO.mkdir(parents=True, exist_ok=True)
    documento.save(ARQUIVO_SAIDA)
    return ARQUIVO_SAIDA


if __name__ == "__main__":
    caminho = gerar_manual()
    print(f"Manual criado em: {caminho}")
